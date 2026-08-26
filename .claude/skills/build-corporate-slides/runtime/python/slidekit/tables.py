from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from .theme import PALETTE
from .typography import set_run

_HEADER_RULE = (PALETTE.text_primary, Pt(1.4))
_ROW_RULE = (PALETTE.line_neutral, Pt(0.75))


def _set_cell_borders(cell, *, top=None, bottom=None, left=None, right=None):
    """セルの罫線を辺ごとに設定する。Noneの辺は罫線なしにする。

    表全体に同じ罫線を四辺一律で引くと、ヘッダーと本文の境界も列の区切りも
    同じ強さになり、構造が読みにくく・のっぺりして見える（縦線は本文の
    余白で区切りを示せば十分で、design-system.mdも「縦線は必要最小限」と
    定めている）。ヘッダー下だけ強い罫線、本文行は薄い横線のみ、縦線は
    引かないという役割分担にするため、辺ごとに個別指定できるようにする。
    """
    properties = cell._tc.get_or_add_tcPr()
    for tag, spec in (("a:lnL", left), ("a:lnR", right),
                      ("a:lnT", top), ("a:lnB", bottom)):
        existing = properties.find(qn(tag))
        if existing is not None:
            properties.remove(existing)
        line = OxmlElement(tag)
        if spec is None:
            line.append(OxmlElement("a:noFill"))
        else:
            color, width = spec
            line.set("w", str(int(width)))
            fill = OxmlElement("a:solidFill")
            value = OxmlElement("a:srgbClr")
            value.set("val", str(color))
            fill.append(value)
            line.append(fill)
        properties.append(line)


def add_data_table(slide, region, columns, rows, *, highlight_key="_highlight"):
    typography = slide._slidekit_typography
    shape = slide.shapes.add_table(len(rows) + 1, len(columns),
                                   region.x, region.y, region.w, region.h)
    table = shape.table
    total_weight = sum(column.get("weight", 1) for column in columns)
    for index, column in enumerate(columns):
        table.columns[index].width = int(
            region.w * column.get("weight", 1) / total_weight)

    for row_index in range(len(rows) + 1):
        table.rows[row_index].height = int(region.h / (len(rows) + 1))
        row_data = None if row_index == 0 else rows[row_index - 1]
        highlighted = bool(row_data and row_data.get(highlight_key))
        for col_index, column in enumerate(columns):
            cell = table.cell(row_index, col_index)
            cell.margin_left = cell.margin_right = Inches(0.16)
            cell.margin_top = cell.margin_bottom = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = PALETTE.surface_subtle
                value = column.get("label", column.get("key", ""))
            else:
                cell.fill.fore_color.rgb = (
                    PALETTE.surface_brand_soft if highlighted else PALETTE.surface_base)
                value = str(row_data.get(column["key"], ""))
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = (PP_ALIGN.CENTER if column.get("align") == "center"
                                   else PP_ALIGN.LEFT)
            run = paragraph.add_run(); run.text = value
            set_run(run,
                    size=typography.body,
                    color=(PALETTE.blue if highlighted and row_index else
                           PALETTE.text_primary),
                    bold=(row_index == 0 or highlighted),
                    font=typography.body_font)
            # 縦線は引かない（列の区切りは余白で示す）。ヘッダー下だけ強い
            # 罫線を引いて本文と分離し、本文行は薄い横線だけで区切る。
            _set_cell_borders(cell, bottom=_HEADER_RULE if row_index == 0 else _ROW_RULE)
    return shape
