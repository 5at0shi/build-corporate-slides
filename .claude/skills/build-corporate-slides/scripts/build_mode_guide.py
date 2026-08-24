"""モード解説（business/dense/large-room）を生成する。

.claude/skills/build-corporate-slides/user-guide/mode-guide.pptx として保存する。
実際のtypographyトークン（theme.py）から数値を読み込むため、スケールを
変更した後にこのスクリプトを再実行すれば自動的に最新化される。

使い方（ワークスペース直下から）:
    ./.venv/bin/python .claude/skills/build-corporate-slides/scripts/build_mode_guide.py
    ./.venv/bin/python .claude/skills/build-corporate-slides/scripts/render_and_check.py \\
        .claude/skills/build-corporate-slides/user-guide/mode-guide.pptx \\
        --pdf .claude/skills/build-corporate-slides/user-guide/mode-guide.pdf
"""
import sys
from pathlib import Path

SKILL = Path(__file__).resolve().parents[1]
WORKSPACE_ROOT = Path.cwd()
sys.path.insert(0, str(SKILL / "runtime" / "python"))

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from slidekit import DeckBuilder
from slidekit.components import add_background_zone, add_key_message, add_section_lead
from slidekit.layout import content_region
from slidekit.theme import PALETTE, TYPE_BUSINESS, TYPE_DENSE, TYPE_LARGE_ROOM
from slidekit.typography import add_paragraph_textbox, add_textbox

MODES = [
    ("business", TYPE_BUSINESS, "標準。個人PC閲覧・事前配布・オンライン会議の画面共有。",
     "通常の社内資料はこのモードで作る。deck.mode は既定でこれ。"),
    ("dense", TYPE_DENSE, "情報量が多いページ（表、比較、条件一覧）。",
     "資料全体をdenseにはしない。該当ページだけslide個別のdensityで指定する。"
     "文字を縮める前に、重複削除・統合・表化・ページ分割を検討する。"),
    ("large-room", TYPE_LARGE_ROOM, "大会議室・講演など遠距離投影。",
     "ユーザーが明示した場合だけ使う。通常の社内資料では使わない。"),
]


def build_mode_page(builder, page, name, type_, when, note):
    slide, area = builder.add_slide(f"モード解説: {name}", page=page)
    left, right = area.columns([1, 1], gap="wide")

    add_section_lead(slide, left.x, left.y, left.w, "いつ使うか")
    add_paragraph_textbox(slide, left.x, left.y + Inches(0.55), left.w, Inches(1.2), [
        {"segments": [(when, {"size": Pt(14), "color": PALETTE.text_primary,
                              "bold": True, "font": "Hiragino Sans"})]}])
    add_paragraph_textbox(slide, left.x, left.y + Inches(1.5), left.w, Inches(2.0), [
        {"segments": [(note, {"size": Pt(12), "color": PALETTE.text_secondary,
                              "font": "Hiragino Sans"})], "line_spacing": 1.3}])

    add_background_zone(slide, right.x, right.y, right.w, right.h,
                        tone="neutral", rounded=True)
    inner = right.inset(Inches(0.35), Inches(0.3))
    add_textbox(slide, inner.x, inner.y, inner.w, Inches(0.3),
                "実寸サンプル（このモードの実際のサイズ）",
                size=Pt(11), color=PALETTE.text_secondary, bold=True)
    y = inner.y + Inches(0.45)
    for label, size in [("Title", type_.title), ("Section", type_.section),
                        ("Body", type_.body), ("Small / Note", type_.small)]:
        add_textbox(slide, inner.x, y, inner.w, Inches(0.15), f"{label}  {size.pt:g}pt",
                    size=Pt(9), color=PALETTE.blue, bold=True)
        y += Inches(0.22)
        sample = "サンプルテキスト Sample 123" if label != "Title" else "サンプル見出し"
        add_textbox(slide, inner.x, y, inner.w, Inches(0.6), sample,
                    size=size, color=PALETTE.text_primary,
                    bold=(label in ("Title", "Section")), font="Hiragino Sans")
        y += Inches(0.75) if label == "Title" else Inches(0.5)

    add_key_message(
        slide, area.x, Inches(6.6), area.w,
        f"title {type_.title.pt:g}pt / section {type_.section.pt:g}pt / "
        f"body {type_.body.pt:g}pt / note {type_.small.pt:g}pt",
        style="subtle")


builder = DeckBuilder.from_workspace(WORKSPACE_ROOT)
builder.add_cover(
    "モード解説", subtitle="business / dense / large-room の使い分け",
    eyebrow="BUILD-CORPORATE-SLIDES")
for index, (name, type_, when, note) in enumerate(MODES, start=2):
    build_mode_page(builder, index, name, type_, when, note)

output = builder.save(SKILL / "user-guide" / "mode-guide.pptx")
print("PPTX:", output)
